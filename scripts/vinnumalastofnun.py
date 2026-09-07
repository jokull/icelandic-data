"""Vinnumálastofnun (Directorate of Labour) — mælaborð + Excel fetcher.

Two sources combined:
  1. Power BI dashboard "Gagnvirk tölfræði Vinnumálastofnunar"
  2. Excel workbook "Helstu talnagögn um atvinnuleysi"

Usage:
    uv run python scripts/vinnumalastofnun.py fetch     # both
    uv run python scripts/vinnumalastofnun.py excel     # Excel only
    uv run python scripts/vinnumalastofnun.py powerbi   # Power BI only
"""
import argparse
import asyncio
import base64
import json
import re
import hashlib
import io
import math
from datetime import datetime
import sys
from pathlib import Path

import httpx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

RAW_DIR = Path(__file__).parent.parent / "data" / "raw" / "vinnumalastofnun"

POWERBI_TENANT = "764a306d-0a68-45ad-9f07-6f1804447cd4"
POWERBI_REPORT_KEY = "e74521bb-e501-4b02-8aa2-08a8bb84d087"
POWERBI_PAGE = "ReportSection7e7dca64570c18a74eb9"

# Resolve the current upload every time: old Contentful assets keep serving 200.
LANDING = "https://island.is/s/vinnumalastofnun/maelabord-og-toelulegar-upplysingar"


def _embed_url() -> str:
    payload = {"k": POWERBI_REPORT_KEY, "t": POWERBI_TENANT, "c": 8}
    token = base64.b64encode(json.dumps(payload, separators=(",", ":")).encode()).decode().rstrip("=")
    return f"https://app.powerbi.com/view?r={token}&pageName={POWERBI_PAGE}"


def resolve_excel(html):
    urls = set(re.findall(r'https://assets\.ctfassets\.net/[^\s"<>]+?Talnagogn_atvinnuleysi\.xlsm', html.replace('\\/', '/')))
    if len(urls) != 1:
        raise ValueError(f"Expected one current unemployment workbook, found {len(urls)}")
    return urls.pop()


def download_excel():
    with httpx.Client(timeout=90, follow_redirects=True) as client:
        page = client.get(LANDING)
        page.raise_for_status()
        url = resolve_excel(page.text)
        response = client.get(url)
        response.raise_for_status()
    return url, response.content


def parse_records(payload):
    """Source-native G2 monthly rates and G4 month-end duration, without macros."""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    records = []
    try:
        for sheet in ['G2', 'G4']:
            dates = {}; measure = None; group = None
            for cells in book[sheet].iter_rows(values_only=True):
                header = {i: v.strftime('%Y-%m') for i, v in enumerate(cells) if isinstance(v, datetime)}
                if header:
                    if len(set(header.values())) != len(header):
                        raise ValueError('Duplicate workbook months')
                    dates = header
                    continue
                labels = [str(v).strip() for v in cells[:3] if v is not None]
                title = ' '.join(labels)
                if sheet == 'G2':
                    if title == 'Atvinnuleysi %': measure = 'rate'; continue
                    if 'Atvinnulausir, meðalfjöldi' in title: measure = None; continue
                    if not measure or not cells[0]: continue
                    group = str(cells[0]).strip(); label = group
                else:
                    if 'Atvinnulausir eftir lengd á skrá' in title: measure = 'duration'; continue
                    if measure != 'duration': continue
                    if cells[0] is not None: group = str(cells[0]).strip()
                    label = str(cells[1]).strip() if cells[1] is not None else ''
                    if label not in ['Allir', 'Karlar', 'Konur', '0-6 mánuðir', '6-12 mánuðir', '+12 mánuðir']: continue
                for column, date in dates.items():
                    value = cells[column] if column < len(cells) else None
                    if value is None: continue
                    if not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0:
                        raise ValueError(f'Invalid {sheet} value at {date}: {value!r}')
                    if measure == 'rate' and value > 1: raise ValueError('Expected fractional rates')
                    records.append(dict(date=date, measure=measure, group=group, label=label,
                                        value=value * 100 if measure == 'rate' else value,
                                        unit='%' if measure == 'rate' else 'persons'))
    finally:
        book.close()
    keys = [(r['date'], r['measure'], r['group'], r['label']) for r in records]
    if not records or len(set(keys)) != len(keys): raise ValueError('Empty or duplicate records')
    for measure in ['rate', 'duration']:
        if not any(r['measure'] == measure for r in records): raise ValueError(f'Missing {measure}')
    national = [r for r in records if r['measure'] == 'duration' and r['group'] == 'Landið allt']
    latest = max(r['date'] for r in national)
    values = {r['label']: r['value'] for r in national if r['date'] == latest}
    if set(values) != {'Allir', '0-6 mánuðir', '6-12 mánuðir', '+12 mánuðir'}:
        raise ValueError('Incomplete latest national duration data')
    if abs(values['Allir'] - sum(values[k] for k in values if k != 'Allir')) > .01:
        raise ValueError('Duration bands do not reconcile')
    return sorted(records, key=lambda r: (r['date'], r['measure'], r['group'], r['label']))


def cmd_excel(args=None):
    url, payload = download_excel()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / 'Talnagogn_atvinnuleysi.xlsm').write_bytes(payload)
    print(url, file=sys.stderr)


def cmd_records(args):
    url, payload = download_excel()
    rows = parse_records(payload)
    digest = hashlib.sha256(payload).hexdigest()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / f'{digest}.xlsm').write_bytes(payload)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    temp = args.out.with_suffix('.tmp')
    temp.write_text(json.dumps(dict(source=LANDING, source_urls=[url], raw_sha256=digest, rows=rows), ensure_ascii=False, allow_nan=False), encoding='utf-8')
    temp.replace(args.out)


async def _scrape_powerbi() -> list[dict]:
    from playwright.async_api import async_playwright

    url = _embed_url()
    results: list[dict] = []
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        try:
            page = await browser.new_page()

            async def on_response(r):
                u = r.url.lower()
                if ("querydata" in u or "executequeries" in u) and r.status == 200:
                    try:
                        results.append(await r.json())
                    except Exception:
                        pass

            page.on("response", on_response)
            print(f"Loading {url}", file=sys.stderr)
            await page.goto(url, wait_until="networkidle", timeout=90000)
            await asyncio.sleep(15)
        finally:
            await browser.close()
    return results


def cmd_powerbi(args=None):
    results = asyncio.run(_scrape_powerbi())
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    out = RAW_DIR / "powerbi.json"
    out.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  captured {len(results)} query responses → {out}", file=sys.stderr)


def cmd_fetch(args=None):
    cmd_excel()
    cmd_powerbi()


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name, fn in [("fetch", cmd_fetch), ("excel", cmd_excel), ("powerbi", cmd_powerbi)]:
        p = sub.add_parser(name)
        p.set_defaults(func=fn)
    p = sub.add_parser('records')
    p.add_argument('--out', type=Path, required=True)
    p.set_defaults(func=cmd_records)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
