#!/usr/bin/env python3
"""Fetch CBI household net lending (banks/pensions separately) or policy rates.

JSON output retains coverage. Mortgage bank values and pension household lending
are published in ISK millions. No assumption that pension lending is bank lending.
"""
import argparse,csv,hashlib,io,json
from datetime import datetime
from pathlib import Path
import httpx
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent
BANK='https://sedlabanki.is/library?itemid=b73e42d6-ba32-4eb3-b39e-1c70d2e45aec'
PENSION='https://fr.sedlabanki.is/sdmx/v2/table/IS2_EXT/LIF_NEW_LOANS_TOTAL/1.0?format=xlsx'

def parse_banks(content):
 s=load_workbook(io.BytesIO(content),data_only=True)['I']
 for row,txt in [(49,'New unindexed'),(87,'New indexed'),(81,'residential mortgage'),(82,'residential mortgage'),(119,'residential mortgage'),(120,'residential mortgage')]:
  if txt not in str(s.cell(row,1).value):raise ValueError(f'Bank schema changed at row {row}')
 out=[]
 for col in range(2,s.max_column+1):
  d=s.cell(10,col).value
  if not isinstance(d,datetime):continue
  for kind,rs in [('Indexed',[119,120]),('Non-indexed',[81,82])]:
   vals=[s.cell(r,col).value for r in rs]
   if not all(isinstance(x,(float,int)) for x in vals):raise ValueError('Missing bank value')
   out.append({'date':d.strftime('%Y-%m'),'lender':'Banks','series':kind,'value':sum(vals)})
 return out

def parse_pensions(content):
 s=load_workbook(io.BytesIO(content),data_only=True)['LIF_NEW_LOANS_TOTAL']
 out=[]
 for col in range(3,s.max_column+1):
  d=str(s.cell(3,col).value)
  if len(d)!=7 or d[4]!='-':continue
  vals=[s.cell(r,col).value for r in [4,5,6]]
  if not all(isinstance(x,(int,float)) for x in vals) or abs(vals[0]-vals[1]-vals[2])>2.1:raise ValueError('Pension totals/coverage changed')
  for kind,r in [('Indexed',5),('Non-indexed',6)]:out.append({'date':d,'lender':'Pension funds','series':kind,'value':s.cell(r,col).value})
 return out

def main():
 a=argparse.ArgumentParser(description=__doc__);a.add_argument('dataset',choices=['mortgages','rates']);a.add_argument('--out',type=Path);args=a.parse_args()
 raw=ROOT/'data/raw/sedlabanki/snapshots';raw.mkdir(parents=True,exist_ok=True)
 with httpx.Client(timeout=60,follow_redirects=True,headers={'User-Agent':'icelandic-data/1.0'},transport=httpx.HTTPTransport(retries=2)) as c:
  if args.dataset=='rates':
   url='https://sedlabanki.is/xmltimeseries/Default.aspx?DagsFra=2015-01-01&TimeSeriesID=17923&Type=csv';r=c.get(url);r.raise_for_status()
   (raw/(hashlib.sha256(r.content).hexdigest()+'.csv')).write_bytes(r.content)
   rows=[]
   for x in csv.reader(io.StringIO(r.text),delimiter=';'):
    if len(x)!=8 or x[2]!='17923':raise ValueError('Unexpected rates CSV')
    rows.append({'date':datetime.strptime(x[6],'%m/%d/%Y %I:%M:%S %p').strftime('%Y-%m-%d'),'series':'Policy rate','value':float(x[7])})
   result={'source':url,'rows':rows}
  else:
   rows=[]
   for url,parse in [(BANK,parse_banks),(PENSION,parse_pensions)]:
    r=c.post('https://gagnabanki.is/api/download',json={'url':url});r.raise_for_status()
    (raw/(hashlib.sha256(r.content).hexdigest()+'.xlsx')).write_bytes(r.content);rows.extend(parse(r.content))
   result={'source':'https://gagnabanki.is/','source_urls':[BANK,PENSION],'unit':'ISK million','rows':rows}
 text=json.dumps(result,ensure_ascii=False,allow_nan=False)
 if args.out:args.out.parent.mkdir(parents=True,exist_ok=True);args.out.write_text(text,encoding='utf-8')
 else:print(text)
if __name__=='__main__':main()
