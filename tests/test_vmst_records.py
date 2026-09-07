import io
from datetime import datetime
import pytest
from openpyxl import Workbook
from scripts.vinnumalastofnun import parse_records, resolve_excel


def workbook(total=6, missing=False):
    b=Workbook();b.remove(b.active)
    g=b.create_sheet('G2');g.append(['extra heading']);g.append([None,None,datetime(2026,7,1)])
    g.append(['Atvinnuleysi %']);g.append(['Landið allt',None,.04])
    g.append(['Suðurnes alls',None,None]);g.append(['Atvinnulausir, meðalfjöldi á mánuði']);g.append(['Landið allt',None,99])
    g=b.create_sheet('G4');g.append([None,None,None,datetime(2026,7,1)])
    g.append(['Atvinnulausir eftir lengd á skrá']);g.append(['Landið allt','Allir',None,total])
    for label,value in [('0-6 mánuðir',3),('6-12 mánuðir',2),('+12 mánuðir',1)]:
        g.append([None,label,None,None if missing and value==1 else value])
    out=io.BytesIO();b.save(out);return out.getvalue()


def test_current_link_and_ambiguity():
    url='https://assets.ctfassets.net/a/b/c/Talnagogn_atvinnuleysi.xlsm'
    assert resolve_excel(f'"{url}" "{url}"') == url
    with pytest.raises(ValueError): resolve_excel('no workbook')
    with pytest.raises(ValueError): resolve_excel(f'"{url}" "{url.replace("/c/", "/d/")}"')


def test_rate_scaling_missing_cells_and_month_end():
    rows=parse_records(workbook())
    rates=[r for r in rows if r['measure']=='rate']
    assert len(rates)==1 and rates[0]['value']==4
    assert len(rows)==5
    assert all(r['date']=='2026-07' for r in rows)


@pytest.mark.parametrize('kwargs',[{'total':7},{'missing':True}])
def test_incomplete_or_inconsistent_duration_fails(kwargs):
    with pytest.raises(ValueError): parse_records(workbook(**kwargs))
