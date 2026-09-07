from datetime import datetime, date
from openpyxl import Workbook
import pytest
from scripts.sedlabanki_fx import read_reserves


def test_reserves_keep_values_in_their_date_columns(tmp_path):
    p = tmp_path / 'reserves.xlsx'
    w = Workbook(); s = w.active; s.title = 'Sedlabanki'
    s.cell(9,1,'M.kr.');s.cell(60,1,'Gjaldeyrisforði')
    # Deliberate interior gap: dropping blanks would misalign later columns.
    for col,dt,value in [(2,datetime(2026,1,31),100),(4,datetime(2026,3,31),300)]:
        s.cell(9,col,dt);s.cell(60,col,value)
    w.save(p)
    assert read_reserves(p).to_dicts() == [{'date':date(2026,1,31),'reserves_mkr':100.0},{'date':date(2026,3,31),'reserves_mkr':300.0}]
    s.cell(60,1,'Unrelated asset row');w.save(p)
    with pytest.raises(ValueError,match='row or unit'): read_reserves(p)


def test_liquidity_sums_categories_without_double_counting_maturities(tmp_path):
    from scripts.sedlabanki_fx import read_liquidity
    p=tmp_path/'liquidity.xlsx';w=Workbook();a=w.active;a.title='I';d=w.create_sheet('II')
    for s in [a,d]:s.cell(5,1,'M.kr.');s.cell(8,1,datetime(2026,7,31))
    for s,c,label in [(a,2,'Official reserve assets'),(d,2,'Foreign currency loans'),(d,22,'Aggregate short and long'),(d,31,'Other')]:s.cell(7,c,label)
    a.cell(8,2,1000);d.cell(8,2,-30);d.cell(8,22,5);d.cell(8,31,-2)
    d.cell(8,7,-10)  # maturity subset: must not be added again
    w.save(p)
    assert read_liquidity(p).to_dicts()==[{'date':date(2026,7,31),'reserve_assets_mkr':1000.0,'net_drains_12m_mkr':27}]


def test_monthly_preserves_gross_and_net_trades(tmp_path,monkeypatch):
    import polars as pl
    from scripts import sedlabanki_fx as m
    day=date(2026,1,30)
    fx=pl.DataFrame({'date':[day]*4,'series_id':[282,284,285,287],'value_mkr':[1.,100.,20.,30.]})
    monkeypatch.setattr(m,'read_fx_market',lambda _:fx)
    monkeypatch.setattr(m,'read_mid_rate',lambda _:pl.DataFrame({'date':[day],'mid_rate':[140.]}))
    monkeypatch.setattr(m,'read_reserves',lambda _:pl.DataFrame({'date':[day],'reserves_mkr':[1000.]}))
    r=m.build_monthly(None,None,None).to_dicts()[0]
    assert (r['purchases_mkr'],r['sales_mkr'],r['net_purchases_mkr'])==(30.,20.,10.)
    assert r['reserves_mkr']==1000.
