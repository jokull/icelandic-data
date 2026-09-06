import io
from openpyxl import Workbook
import pytest
from scripts.sedlabanki_household import parse_pensions

def test_pension_identity_and_missing_values():
 w=Workbook();s=w.active;s.title='LIF_NEW_LOANS_TOTAL';s.cell(3,3,'2026-07')
 for row,v in [(4,10),(5,8),(6,2)]:s.cell(row,3,v)
 b=io.BytesIO();w.save(b)
 assert sum(x['value'] for x in parse_pensions(b.getvalue()))==10
 s.cell(4,3,100);b=io.BytesIO();w.save(b)
 with pytest.raises(ValueError):parse_pensions(b.getvalue())
